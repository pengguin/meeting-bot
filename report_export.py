import html as html_lib
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from meetingbot_config import REPORT_BODY_FONT, REPORT_HEADING_FONT, get_template_names
from speaker_naming import anonymous_speaker_label


# ============================================================
# DOCX/HTML/MD/PDF 报告导出
# ============================================================

def rgb_from_hex(hex_color: str) -> RGBColor:
    hex_color = hex_color.replace("#", "")
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)

    for attr in ["ascii", "hAnsi", "eastAsia", "cs"]:
        r_fonts.set(qn(f"w:{attr}"), font_name)


def apply_run_style(
    run,
    font_name: str = REPORT_BODY_FONT,
    size: float = 10.5,
    bold: bool = False,
    color: Optional[str] = None,
) -> None:
    set_east_asia_font(run, font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = rgb_from_hex(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9E2F3", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ["top", "left", "bottom", "right"]:
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clear_cell(cell) -> None:
    cell.text = ""


def add_text_to_cell(
    cell,
    text: str,
    bold: bool = False,
    size: float = 10,
    color: Optional[str] = None,
    align: Optional[int] = None,
) -> None:
    clear_cell(cell)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    apply_run_style(run, size=size, bold=bold, color=color)


def set_cell_margins(
    cell,
    top: int = 120,
    start: int = 160,
    bottom: int = 120,
    end: int = 160,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_column_widths(table, widths_cm: List[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def set_paragraph_left_border(
    paragraph,
    color: str = "CBD5E1",
    size: str = "8",
    space: str = "8",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)

    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)


def set_paragraph_bottom_border(
    paragraph,
    color: str = "E5E7EB",
    size: str = "6",
    space: str = "4",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_spacer(doc: Document, points: float = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)
    p.paragraph_format.line_spacing = 1


def add_section_heading(
    doc: Document,
    number: str,
    title: str,
    subtitle: str = "",
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_paragraph_bottom_border(p, color="E5E7EB", size="5", space="3")

    prefix = p.add_run(f"{number}  ")
    apply_run_style(prefix, size=9.5, bold=True, color="2563EB")

    heading = p.add_run(title)
    apply_run_style(heading, font_name=REPORT_HEADING_FONT, size=15.5, bold=True, color="111827")

    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        p2.paragraph_format.line_spacing = 1.08
        sub = p2.add_run(subtitle)
        apply_run_style(sub, size=9, color="6B7280")


def add_inline_label(paragraph, label: str, value: str) -> None:
    label_run = paragraph.add_run(label)
    apply_run_style(label_run, size=8.5, bold=True, color="6B7280")
    value_run = paragraph.add_run(value)
    apply_run_style(value_run, size=9.5, color="1F2937")


def setup_report_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = REPORT_BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), REPORT_BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), REPORT_BODY_FONT)
    normal.font.size = Pt(10.2)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = REPORT_HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), REPORT_HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), REPORT_HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), REPORT_HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), REPORT_HEADING_FONT)

    styles["Heading 1"].font.size = Pt(15.5)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 3"].font.size = Pt(10.5)


def add_report_header_footer(doc: Document, report_title: str) -> None:
    section = doc.sections[0]

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("内容由 AI 生成，仅供参考；正式使用前请人工核对")
    apply_run_style(run, size=8, color="9CA3AF")
    page_run = fp.add_run("  ·  第 ")
    apply_run_style(page_run, size=8, color="9CA3AF")
    add_word_field(fp, "PAGE")
    total_run = fp.add_run(" / ")
    apply_run_style(total_run, size=8, color="9CA3AF")
    add_word_field(fp, "NUMPAGES")
    end_run = fp.add_run(" 页")
    apply_run_style(end_run, size=8, color="9CA3AF")


def add_word_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    set_east_asia_font(run, REPORT_BODY_FONT)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    run.font.size = Pt(8)
    run.font.color.rgb = rgb_from_hex("9CA3AF")


# ============================================================
# 18. DOCX 报告 3.0：报告模块
# ============================================================

def add_report_hero(doc: Document, report: Dict, named: bool) -> None:
    report_title = report.get("report_title", "会议纪要")
    version_name = "实名版" if named else "匿名版"
    meeting_type = get_template_names().get(report.get("meeting_type"), "通用会议纪要")
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.04
    title_run = title.add_run(report_title)
    apply_run_style(title_run, font_name=REPORT_HEADING_FONT, size=22, bold=True, color="111827")

    meta_lines = [
        ("生成时间：", generated_at),
        ("报告版本：", version_name),
        ("会议类型：", meeting_type),
        ("使用说明：", "AI 自动整理生成，适合会后复盘；正式外发前请人工核对"),
    ]
    for idx, (label, value) in enumerate(meta_lines):
        p = doc.add_paragraph()
        set_paragraph_left_border(p, color="CBD5E1", size="8", space="8")
        p.paragraph_format.space_after = Pt(2 if idx < len(meta_lines) - 1 else 0)
        p.paragraph_format.line_spacing = 1.08
        add_inline_label(p, label, value)

    add_spacer(doc, 10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    set_paragraph_bottom_border(p, color="E5E7EB", size="5", space="4")
    run = p.add_run("会议纪要")
    apply_run_style(run, font_name=REPORT_HEADING_FONT, size=17, bold=True, color="111827")

    add_spacer(doc, 2)

def add_metric_strip_to_cell(cell, report: Dict) -> None:
    metrics = report.get("key_metrics", {})
    items = [
        ("结论", metrics.get("conclusion_count", 0), "2563EB"),
        ("行动", metrics.get("action_item_count", 0), "16A34A"),
        ("待确认", metrics.get("open_question_count", 0), "64748B"),
    ]

    table = cell.add_table(rows=1, cols=3)
    set_table_column_widths(table, [2.55, 2.55, 2.55])
    for idx, (label, value, color) in enumerate(items):
        item_cell = table.cell(0, idx)
        set_cell_shading(item_cell, "FFFFFF")
        set_cell_border(item_cell, color="E5E7EB", size="4")
        set_cell_margins(item_cell, top=70, start=80, bottom=70, end=80)
        clear_cell(item_cell)
        p = item_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        value_run = p.add_run(str(value))
        apply_run_style(value_run, size=13.5, bold=True, color=color)
        label_run = p.add_run(f"\n{label}")
        apply_run_style(label_run, size=8.3, color="64748B")


def add_takeaway_block(doc: Document, report: Dict) -> None:
    add_section_heading(
        doc,
        "01",
        "总览",
        "先看结论和摘要，再进入行动项、议题和说话人信息。",
    )

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_column_widths(table, [7.6, 9.3])

    left = table.cell(0, 0)
    set_cell_shading(left, "F8FAFC")
    set_cell_border(left, color="E2E8F0", size="5")
    set_cell_margins(left, top=150, start=170, bottom=150, end=170)
    clear_cell(left)

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    title_run = p.add_run("核心判断")
    apply_run_style(title_run, size=9.3, bold=True, color="2563EB")

    p2 = left.add_paragraph()
    p2.paragraph_format.line_spacing = 1.12
    p2.paragraph_format.space_after = Pt(8)
    r = p2.add_run(report.get("one_sentence_takeaway", "待确认"))
    apply_run_style(r, size=12.6, bold=True, color="111827")

    add_metric_strip_to_cell(left, report)

    right = table.cell(0, 1)
    set_cell_shading(right, "EFF6FF")
    set_cell_border(right, color="CFE0FF", size="5")
    set_cell_margins(right, top=150, start=180, bottom=150, end=180)
    clear_cell(right)

    p3 = right.paragraphs[0]
    p3.paragraph_format.space_after = Pt(4)
    label = p3.add_run("会议摘要")
    apply_run_style(label, size=9.3, bold=True, color="1D4ED8")

    p4 = right.add_paragraph()
    p4.paragraph_format.line_spacing = 1.15
    p4.paragraph_format.space_after = Pt(0)
    r4 = p4.add_run(report.get("executive_summary", "待确认"))
    apply_run_style(r4, size=9.8, color="1F2937")

    add_spacer(doc, 5)


def add_metric_cards(doc: Document, report: Dict) -> None:
    metrics = report.get("key_metrics", {})

    cards = [
        ("核心结论", str(metrics.get("conclusion_count", 0)), "2F6FED", "已提炼"),
        ("行动项", str(metrics.get("action_item_count", 0)), "10B981", "待跟进"),
        ("待确认", str(metrics.get("open_question_count", 0)), "64748B", "需核对"),
    ]

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_column_widths(table, [5.5, 5.5, 5.5])

    for idx, (label, value, color, helper) in enumerate(cards):
        cell = table.cell(0, idx)
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell, color="E5E7EB", size="5")
        set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
        clear_cell(cell)

        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(value)
        apply_run_style(r1, size=20, bold=True, color=color)

        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(label)
        apply_run_style(r2, size=9.5, bold=True, color="111827")

        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(helper)
        apply_run_style(r3, size=8.5, color="6B7280")

    add_spacer(doc, 6)


def add_key_conclusions(doc: Document, report: Dict) -> None:
    conclusions = report.get("key_conclusions", [])
    if not conclusions:
        return

    add_section_heading(doc, "02", "关键要点", "按重要性排列，保留可执行信息。")

    for idx, item in enumerate(conclusions, start=1):
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.keep_with_next = True
        number_run = p1.add_run(f"{idx:02d}  ")
        apply_run_style(number_run, size=9.2, bold=True, color="2563EB")
        r1 = p1.add_run(item.get("title", ""))
        apply_run_style(r1, size=10.8, bold=True, color="111827")

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.72)
        p2.paragraph_format.line_spacing = 1.14
        p2.paragraph_format.space_after = Pt(5)
        r2 = p2.add_run(item.get("detail", ""))
        apply_run_style(r2, size=9.7, color="374151")


def add_action_items(doc: Document, report: Dict) -> None:
    actions = report.get("action_items", [])
    if not actions:
        return

    add_section_heading(doc, "03", "后续安排", "从会议讨论中抽取的待办，方便直接跟进。")

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_column_widths(table, [6.4, 2.4, 2.6, 1.7, 3.9])

    headers = ["行动事项", "负责人", "截止时间", "优先级", "备注"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "F1F5F9")
        set_cell_margins(cell, top=85, start=105, bottom=85, end=105)
        add_text_to_cell(
            cell,
            header,
            bold=True,
            size=9,
            color="334155",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        cell.paragraphs[0].paragraph_format.keep_with_next = True

    priority_colors = {
        "高": ("DBEAFE", "1D4ED8"),
        "中": ("E0F2FE", "0369A1"),
        "低": ("ECFDF5", "047857"),
        "待确认": ("F8FAFC", "475569"),
    }

    for item in actions:
        row = table.add_row()
        priority = item.get("priority", "待确认")
        fill, color = priority_colors.get(priority, priority_colors["待确认"])
        vals = [
            item.get("task", ""),
            item.get("owner", "待确认"),
            item.get("deadline", "待确认"),
            priority,
            item.get("notes", ""),
        ]

        for idx, val in enumerate(vals):
            cell = row.cells[idx]
            set_cell_margins(cell, top=82, start=100, bottom=82, end=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx == 3:
                set_cell_shading(cell, fill)
                add_text_to_cell(
                    cell,
                    val,
                    bold=True,
                    size=8.8,
                    color=color,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
            else:
                add_text_to_cell(cell, val, size=8.8, color="1F2937")

    add_spacer(doc, 4)


def add_discussion_topics(doc: Document, report: Dict) -> None:
    topics = report.get("discussion_topics", [])
    if not topics:
        return

    add_section_heading(doc, "04", "议题复盘", "按主题整理讨论过程，便于回看上下文。")

    for idx, topic in enumerate(topics, start=1):
        p1 = doc.add_paragraph()
        set_paragraph_left_border(p1, color="93C5FD", size="10", space="8")
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(2)
        r0 = p1.add_run(f"{idx:02d}  ")
        apply_run_style(r0, size=9, bold=True, color="2563EB")
        r1 = p1.add_run(topic.get("title", ""))
        apply_run_style(r1, size=10.3, bold=True, color="111827")

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.55)
        p2.paragraph_format.line_spacing = 1.13
        p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(topic.get("summary", ""))
        apply_run_style(r2, size=9.5, color="374151")

        for point in topic.get("points", []):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(point)
            apply_run_style(r, size=9.1, color="4B5563")

        add_spacer(doc, 3)


def add_open_questions(doc: Document, report: Dict) -> None:
    questions = report.get("open_questions", [])
    if not questions:
        return

    add_section_heading(doc, "05", "待确认问题", "需要进一步确认、补充材料或后续决策的事项。")

    for idx, item in enumerate(questions, start=1):
        p1 = doc.add_paragraph()
        set_paragraph_left_border(p1, color="94A3B8", size="10", space="8")
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(f"Q{idx}  {item.get('question', '')}")
        apply_run_style(r1, size=10.2, bold=True, color="334155")

        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.55)
        p2.paragraph_format.line_spacing = 1.1
        p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(f"重要性：{item.get('why_it_matters', '')}")
        apply_run_style(r2, size=9.2, color="475569")

        add_spacer(doc, 2)


def add_speaker_insights(doc: Document, report: Dict) -> None:
    speakers = report.get("speaker_insights", [])
    if not speakers:
        return

    add_section_heading(doc, "06", "发言人观点", "保留不同角色的主要观点，方便会后对齐。")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_column_widths(table, [2.8, 2.8, 11.3])

    headers = ["发言人", "角色", "主要观点"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "F1F5F9")
        set_cell_margins(cell, top=85, start=105, bottom=85, end=105)
        add_text_to_cell(
            cell,
            h,
            bold=True,
            size=9,
            color="334155",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for speaker in speakers:
        row = table.add_row()
        for cell in row.cells:
            set_cell_margins(cell, top=85, start=105, bottom=85, end=105)

        add_text_to_cell(row.cells[0], speaker.get("speaker", ""), size=9, color="1F2937")
        add_text_to_cell(row.cells[1], speaker.get("role", "待确认"), size=9, color="4B5563")

        clear_cell(row.cells[2])
        views = speaker.get("main_views", [])
        if not views:
            views = ["待确认"]

        for idx, view in enumerate(views):
            p = row.cells[2].paragraphs[0] if idx == 0 else row.cells[2].add_paragraph()
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(f"• {view}")
            apply_run_style(r, size=9, color="374151")

    add_spacer(doc, 4)


def add_speaker_mapping_table(
    doc: Document,
    speaker_map: Dict[str, str],
    named: bool,
) -> None:
    add_section_heading(doc, "07", "说话人标注", "统一使用匿名标签或已标注姓名展示。")

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_column_widths(table, [7.0, 9.9])

    headers = ["匿名标签", "当前显示名称"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "F1F5F9")
        set_cell_margins(cell, top=75, start=105, bottom=75, end=105)
        add_text_to_cell(
            cell,
            h,
            bold=True,
            size=9,
            color="374151",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for raw, name in speaker_map.items():
        row = table.add_row()
        for cell in row.cells:
            set_cell_margins(cell, top=70, start=105, bottom=70, end=105)
        add_text_to_cell(
            row.cells[0],
            anonymous_speaker_label(raw, fallback=name),
            size=9,
            color="4B5563",
        )
        add_text_to_cell(row.cells[1], name, size=9, color="111827")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    if named:
        text = "说明：本报告已根据用户补充信息完成说话人身份更新。"
    else:
        text = "说明：当前为匿名版，可在飞书中继续回复“说话人1=姓名”等指令更新。"
    r = note.add_run(text)
    apply_run_style(r, size=8.8, color="6B7280")

    add_spacer(doc, 3)


def add_report_notes(doc: Document, report: Dict) -> None:
    notes = report.get("report_notes", [])
    if not notes:
        return

    add_section_heading(doc, "08", "说明")
    for note in notes:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(note)
        apply_run_style(r, size=9, color="6B7280")


def safe_filename_component(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", plain_text(value))
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    return cleaned or "会议纪要"


def build_formal_minutes_output_path(
    session_path: Path,
    report_title: str,
    version_name: str,
) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_title = safe_filename_component(report_title)
    base_name = f"{safe_title}_{version_name}_{timestamp}"
    output_path = session_path / f"{base_name}.docx"

    counter = 2
    while output_path.exists():
        output_path = session_path / f"{base_name}_{counter:02d}.docx"
        counter += 1

    return output_path


def generate_formal_minutes_docx(
    report: Dict,
    classification: Dict,
    session_path: Path,
    named: bool,
    speaker_map: Dict[str, str],
) -> Path:
    version_name = "实名版" if named else "匿名版"
    output_path = build_formal_minutes_output_path(
        session_path=session_path,
        report_title=report.get("report_title", "会议纪要"),
        version_name=version_name,
    )

    doc = Document()
    setup_report_document(doc)
    add_report_header_footer(
        doc,
        report_title=report.get("report_title", "会议纪要"),
    )

    add_report_hero(doc, report, named=named)
    add_takeaway_block(doc, report)
    add_key_conclusions(doc, report)
    add_action_items(doc, report)
    add_discussion_topics(doc, report)
    add_open_questions(doc, report)
    add_speaker_insights(doc, report)
    add_speaker_mapping_table(doc, speaker_map, named=named)
    add_report_notes(doc, report)

    doc.save(output_path)
    return output_path


def plain_text(value) -> str:
    return str(value or "").strip()


def md_cell(value) -> str:
    return plain_text(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def html_escape(value) -> str:
    return html_lib.escape(plain_text(value), quote=True)


def html_paragraphs(value) -> str:
    lines = [line.strip() for line in html_escape(value).splitlines() if line.strip()]
    return "".join(f"<p>{line}</p>" for line in lines)


def report_base_meta(report: Dict, named: bool) -> Dict[str, str]:
    version_name = "实名版" if named else "匿名版"
    meeting_type = get_template_names().get(report.get("meeting_type"), "通用会议纪要")
    return {
        "title": report.get("report_title", "会议纪要"),
        "version_name": version_name,
        "meeting_type": meeting_type,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def generate_formal_minutes_markdown(
    report: Dict,
    output_path: Path,
    named: bool,
    speaker_map: Dict[str, str],
) -> Path:
    meta = report_base_meta(report, named)
    metrics = report.get("key_metrics", {})
    lines = [
        f"# {plain_text(meta['title'])}",
        "",
        f"- 生成时间：{meta['generated_at']}",
        f"- 报告版本：{meta['version_name']}",
        f"- 会议类型：{meta['meeting_type']}",
        "- 使用说明：AI 自动整理生成，适合会后复盘；正式外发前请人工核对",
        "",
        "## 会议纪要",
        "",
        "## 01 总览",
        "",
        f"**核心判断：** {plain_text(report.get('one_sentence_takeaway', '待确认'))}",
        "",
        plain_text(report.get("executive_summary", "待确认")),
        "",
        "| 结论 | 行动 | 待确认 |",
        "| --- | --- | --- |",
        f"| {metrics.get('conclusion_count', 0)} | {metrics.get('action_item_count', 0)} | {metrics.get('open_question_count', 0)} |",
        "",
    ]

    conclusions = report.get("key_conclusions", [])
    if conclusions:
        lines.extend(["## 02 关键要点", ""])
        for idx, item in enumerate(conclusions, start=1):
            lines.extend([f"### {idx:02d} {plain_text(item.get('title'))}", "", plain_text(item.get("detail")), ""])

    actions = report.get("action_items", [])
    if actions:
        lines.extend([
            "## 03 后续安排",
            "",
            "| 行动事项 | 负责人 | 截止时间 | 优先级 | 备注 |",
            "| --- | --- | --- | --- | --- |",
        ])
        for item in actions:
            row = [
                md_cell(item.get("task")),
                md_cell(item.get("owner", "待确认")),
                md_cell(item.get("deadline", "待确认")),
                md_cell(item.get("priority", "待确认")),
                md_cell(item.get("notes")),
            ]
            lines.append("| " + " | ".join(row) + " |")
        lines.append("")

    topics = report.get("discussion_topics", [])
    if topics:
        lines.extend(["## 04 议题复盘", ""])
        for idx, topic in enumerate(topics, start=1):
            lines.extend([f"### {idx:02d} {plain_text(topic.get('title'))}", "", plain_text(topic.get("summary")), ""])
            for point in topic.get("points", []):
                lines.append(f"- {plain_text(point)}")
            lines.append("")

    questions = report.get("open_questions", [])
    if questions:
        lines.extend(["## 05 待确认问题", ""])
        for idx, item in enumerate(questions, start=1):
            lines.extend([
                f"### Q{idx} {plain_text(item.get('question'))}",
                "",
                f"重要性：{plain_text(item.get('why_it_matters'))}",
                "",
            ])

    speakers = report.get("speaker_insights", [])
    if speakers:
        lines.extend(["## 06 发言人观点", "", "| 发言人 | 角色 | 主要观点 |", "| --- | --- | --- |"])
        for speaker in speakers:
            views = "<br>".join(plain_text(v) for v in speaker.get("main_views", []))
            lines.append(
                "| "
                + " | ".join([md_cell(speaker.get("speaker")), md_cell(speaker.get("role", "待确认")), md_cell(views)])
                + " |"
            )
        lines.append("")

    if speaker_map:
        lines.extend(["## 07 说话人标注", "", "| 匿名标签 | 当前显示名称 |", "| --- | --- |"])
        for raw, name in speaker_map.items():
            lines.append(f"| {md_cell(anonymous_speaker_label(raw, fallback=name))} | {md_cell(name)} |")
        lines.append("")

    notes = report.get("report_notes", [])
    if notes:
        lines.extend(["## 08 说明", ""])
        for note in notes:
            lines.append(f"- {plain_text(note)}")
        lines.append("")

    output_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return output_path


def generate_formal_minutes_html(
    report: Dict,
    output_path: Path,
    named: bool,
    speaker_map: Dict[str, str],
) -> Path:
    meta = report_base_meta(report, named)
    metrics = report.get("key_metrics", {})

    def section_html(title: str, body: str, subtitle: str = "") -> str:
        subtitle_part = f'<p class="section-subtitle">{html_escape(subtitle)}</p>' if subtitle else ""
        number, _, heading = title.partition(" ")
        if number.isdigit() and heading:
            title_html = (
                f'<span class="section-number">{html_escape(number)}</span>'
                f"{html_escape(heading)}"
            )
        else:
            title_html = html_escape(title)
        return f"<section><h2>{title_html}</h2>{subtitle_part}{body}</section>"

    sections: List[str] = []
    overview = f"""
    <div class="overview">
      <div>
        <p class="label">核心判断</p>
        <p class="takeaway">{html_escape(report.get("one_sentence_takeaway", "待确认"))}</p>
        <div class="metrics">
          <div class="metric"><strong>{metrics.get("conclusion_count", 0)}</strong><span>结论</span></div>
          <div class="metric"><strong>{metrics.get("action_item_count", 0)}</strong><span>行动</span></div>
          <div class="metric"><strong>{metrics.get("open_question_count", 0)}</strong><span>待确认</span></div>
        </div>
      </div>
      <div>
        <p class="label">会议摘要</p>
        {html_paragraphs(report.get("executive_summary", "待确认"))}
      </div>
    </div>
    """
    sections.append(section_html("01 总览", overview, "先看结论和摘要，再进入行动项、议题和说话人信息。"))

    conclusions_html = ""
    for idx, item in enumerate(report.get("key_conclusions", []), start=1):
        conclusions_html += (
            '<article class="line-item">'
            f"<h3><span>{idx:02d}</span>{html_escape(item.get('title'))}</h3>"
            f"{html_paragraphs(item.get('detail'))}"
            "</article>"
        )
    if conclusions_html:
        sections.append(section_html("02 关键要点", conclusions_html, "按重要性排列，保留可执行信息。"))

    actions = report.get("action_items", [])
    if actions:
        rows = ""
        for item in actions:
            rows += (
                "<tr>"
                f"<td>{html_escape(item.get('task'))}</td>"
                f"<td>{html_escape(item.get('owner', '待确认'))}</td>"
                f"<td>{html_escape(item.get('deadline', '待确认'))}</td>"
                f"<td><span class=\"tag\">{html_escape(item.get('priority', '待确认'))}</span></td>"
                f"<td>{html_escape(item.get('notes'))}</td>"
                "</tr>"
            )
        table = (
            '<div class="table-wrap"><table><thead><tr>'
            "<th>行动事项</th><th>负责人</th><th>截止时间</th><th>优先级</th><th>备注</th>"
            f"</tr></thead><tbody>{rows}</tbody></table></div>"
        )
        sections.append(section_html("03 后续安排", table, "从会议讨论中抽取的待办，方便直接跟进。"))

    topics_html = ""
    for idx, topic in enumerate(report.get("discussion_topics", []), start=1):
        points = "".join(f"<li>{html_escape(point)}</li>" for point in topic.get("points", []))
        topics_html += (
            '<article class="line-item">'
            f"<h3><span>{idx:02d}</span>{html_escape(topic.get('title'))}</h3>"
            f"{html_paragraphs(topic.get('summary'))}<ul>{points}</ul>"
            "</article>"
        )
    if topics_html:
        sections.append(section_html("04 议题复盘", topics_html, "按主题整理讨论过程，便于回看上下文。"))

    questions_html = ""
    for idx, item in enumerate(report.get("open_questions", []), start=1):
        questions_html += (
            '<article class="line-item muted">'
            f"<h3><span>Q{idx}</span>{html_escape(item.get('question'))}</h3>"
            f"<p>重要性：{html_escape(item.get('why_it_matters'))}</p>"
            "</article>"
        )
    if questions_html:
        sections.append(section_html("05 待确认问题", questions_html, "需要进一步确认、补充材料或后续决策的事项。"))

    speakers = report.get("speaker_insights", [])
    if speakers:
        rows = ""
        for speaker in speakers:
            views = "".join(f"<li>{html_escape(v)}</li>" for v in speaker.get("main_views", []))
            rows += (
                "<tr>"
                f"<td>{html_escape(speaker.get('speaker'))}</td>"
                f"<td>{html_escape(speaker.get('role', '待确认'))}</td>"
                f"<td><ul>{views}</ul></td>"
                "</tr>"
            )
        table = (
            '<div class="table-wrap"><table><thead><tr>'
            "<th>发言人</th><th>角色</th><th>主要观点</th>"
            f"</tr></thead><tbody>{rows}</tbody></table></div>"
        )
        sections.append(section_html("06 发言人观点", table, "保留不同角色的主要观点，方便会后对齐。"))

    if speaker_map:
        rows = "".join(
            f"<tr><td>{html_escape(anonymous_speaker_label(raw, fallback=name))}</td><td>{html_escape(name)}</td></tr>"
            for raw, name in speaker_map.items()
        )
        table = (
            '<div class="table-wrap"><table><thead><tr>'
            "<th>匿名标签</th><th>当前显示名称</th>"
            f"</tr></thead><tbody>{rows}</tbody></table></div>"
        )
        sections.append(section_html("07 说话人标注", table, "统一使用匿名标签或已标注姓名展示。"))

    notes = report.get("report_notes", [])
    if notes:
        notes_html = "<ul>" + "".join(f"<li>{html_escape(note)}</li>" for note in notes) + "</ul>"
        sections.append(section_html("08 说明", notes_html))

    html = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html_escape(meta['title'])}</title>
  <style>
    :root {{
      --text: #111827;
      --muted: #64748b;
      --line: #e2e8f0;
      --soft: #f8fafc;
      --soft-blue: #eff6ff;
      --blue: #2563eb;
      --blue-dark: #1d4ed8;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      color: var(--text);
      background: #f8fafc;
      font-family: "PingFang SC", "Microsoft YaHei", "Source Han Sans SC", "Noto Sans CJK SC", Arial, sans-serif;
      line-height: 1.62;
    }}
    main {{
      width: min(960px, calc(100vw - 32px));
      margin: 40px auto;
      padding: 44px 52px;
      background: #fff;
      border: 1px solid var(--line);
      border-radius: 18px;
      box-shadow: 0 18px 48px rgba(15, 23, 42, 0.08);
    }}
    h1 {{ margin: 0 0 20px; font-size: 34px; line-height: 1.25; letter-spacing: 0; }}
    .meta {{ margin: 0 0 34px; padding-left: 16px; border-left: 3px solid #cbd5e1; color: #334155; font-size: 14px; }}
    .meta p {{ margin: 4px 0; }}
    .chapter {{ margin: 12px 0 30px; padding-bottom: 12px; border-bottom: 1px solid var(--line); font-size: 26px; }}
    section {{ margin: 30px 0; }}
    h2 {{ margin: 0 0 10px; padding-bottom: 9px; border-bottom: 1px solid var(--line); font-size: 24px; line-height: 1.3; }}
    .section-number {{ margin-right: 8px; color: var(--blue); }}
    .section-subtitle {{ margin: 0 0 16px; color: var(--muted); }}
    .overview {{ display: grid; grid-template-columns: minmax(0, 1fr) minmax(0, 1.05fr); border: 1px solid #cfe0ff; background: var(--soft-blue); }}
    .overview > div {{ padding: 18px; }}
    .overview > div:first-child {{ background: var(--soft); border-right: 1px solid #cfe0ff; }}
    .label {{ color: var(--blue); font-size: 14px; font-weight: 700; }}
    .takeaway {{ font-size: 20px; font-weight: 700; line-height: 1.55; }}
    .metrics {{ display: grid; grid-template-columns: repeat(3, 1fr); border: 1px solid var(--line); background: #fff; }}
    .metric {{ padding: 10px; text-align: center; border-right: 1px solid var(--line); }}
    .metric:last-child {{ border-right: 0; }}
    .metric strong {{ display: block; color: var(--blue); font-size: 22px; }}
    .metric span {{ color: var(--muted); font-size: 13px; }}
    .line-item {{ margin: 16px 0; padding-left: 14px; border-left: 3px solid #93c5fd; }}
    .line-item.muted {{ border-left-color: #94a3b8; }}
    h3 {{ margin: 0 0 8px; font-size: 17px; }}
    h3 span {{ margin-right: 8px; color: var(--blue); font-size: 14px; }}
    p {{ margin: 7px 0; }}
    ul {{ margin: 8px 0 0 20px; padding: 0; }}
    li {{ margin: 4px 0; }}
    .table-wrap {{ overflow-x: auto; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 14px; }}
    th, td {{ padding: 10px 12px; border: 1px solid var(--line); vertical-align: top; }}
    th {{ background: #f1f5f9; color: #334155; text-align: left; }}
    .tag {{ display: inline-block; min-width: 42px; padding: 2px 9px; border-radius: 999px; background: #dbeafe; color: var(--blue-dark); text-align: center; font-weight: 700; }}
    footer {{ margin-top: 42px; padding-top: 14px; border-top: 1px solid var(--line); color: #94a3b8; font-size: 13px; text-align: center; }}
    @media print {{
      @page {{
        margin: 16mm 14mm 18mm;
        @bottom-center {{
          content: "第 " counter(page) " / " counter(pages) " 页";
          color: #94a3b8;
          font-size: 10px;
          font-family: "PingFang SC", "Microsoft YaHei", "Source Han Sans SC", "Noto Sans CJK SC", Arial, sans-serif;
        }}
      }}
      body {{ background: #fff; }}
      main {{ width: auto; margin: 0; padding: 28px; border: 0; box-shadow: none; }}
      .overview, .line-item, table {{ break-inside: avoid; }}
    }}
    @media (max-width: 720px) {{
      main {{ padding: 28px 20px; }}
      h1 {{ font-size: 28px; }}
      .overview {{ grid-template-columns: 1fr; }}
      .overview > div:first-child {{ border-right: 0; border-bottom: 1px solid #cfe0ff; }}
    }}
  </style>
</head>
<body>
<main>
  <h1>{html_escape(meta['title'])}</h1>
  <div class="meta">
    <p>生成时间：{html_escape(meta['generated_at'])}</p>
    <p>报告版本：{html_escape(meta['version_name'])}</p>
    <p>会议类型：{html_escape(meta['meeting_type'])}</p>
    <p>使用说明：AI 自动整理生成，适合会后复盘；正式外发前请人工核对</p>
  </div>
  <h1 class="chapter">会议纪要</h1>
  {''.join(sections)}
  <footer>内容由 AI 生成，仅供参考；正式使用前请人工核对</footer>
</main>
</body>
</html>
"""
    output_path.write_text(html, encoding="utf-8")
    return output_path


# ============================================================
# 19. DOCX 转 PDF
# ============================================================

def find_soffice_command() -> str:
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]

    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate

    raise RuntimeError(
        "未找到 LibreOffice 的 soffice 命令。"
        "请确认已安装 LibreOffice，或将 soffice 加入 PATH。"
    )


def convert_docx_to_pdf(docx_path: Path) -> Path:
    soffice = find_soffice_command()

    cmd = [
        soffice,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(docx_path.parent),
        str(docx_path),
    ]

    result = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        timeout=600,
    )

    if result.returncode != 0:
        raise RuntimeError(
            "DOCX 转 PDF 失败。\n"
            f"STDOUT:\n{result.stdout}\n"
            f"STDERR:\n{result.stderr}"
        )

    pdf_path = docx_path.with_suffix(".pdf")
    if not pdf_path.exists():
        raise RuntimeError("LibreOffice 未生成 PDF 文件")

    return pdf_path
